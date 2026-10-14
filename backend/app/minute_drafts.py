from __future__ import annotations

import copy
import hashlib
import io
import json
import re
import sqlite3
import uuid
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .minutes import (_date_from_text, _participants_from_text, _project_from_text,
                      _project_responsible_from_text, canonical, normalize_text)
from .minutes_deterministic import extract_explicit
from .minutes_enrichment import build_activities
from .minutes_semantic import extract_rule_candidates


HEADERS = (
    "Actividad *", "Área_Responsable", "Responsable *", "Tipo_de_seguimiento *", "Fecha_Objetivo",
    "Proyecto_o_Frente", "Contexto_o_Resultado", "Comentarios",
)
EDITABLE_FIELDS = {
    "actividad", "area_responsable", "responsable", "responsable_account_id", "tipo_de_seguimiento",
    "tipo_de_seguimiento_id", "fecha_objetivo", "proyecto_o_frente",
    "contexto_o_resultado", "comentarios", "ignored_suggestions",
    "responsable_resolution_status", "area_source",
}
PROJECT_RESPONSIBLE_EDITABLE_FIELDS = {
    "project_responsible", "project_responsible_account_id",
    "project_responsible_resolution_status", "project_responsible_lookup_query",
}
RESPONSABLE_STATES = {"UNRESOLVED", "RESOLVING", "EXPLICIT_MINUTES", "AUTO_RESOLVED",
                      "VALIDATED_JIRA", "MULTIPLE_MATCHES", "NOT_FOUND", "UNAVAILABLE", "NEEDS_REVIEW"}


class MinuteDraftError(Exception):
    def __init__(self, message: str, status_code: int = 422, code: str = "MINUTE_DRAFT_ERROR"):
        super().__init__(message)
        self.status_code = status_code
        self.code = code


def _now() -> str:
    return datetime.now(UTC).isoformat().replace("+00:00", "Z")


def _clean(value: Any, limit: int) -> str | None:
    if value is None:
        return None
    result = re.sub(r"\s+", " ", str(value)).strip()
    return result[:limit] or None


def _responsible_from_evidence(evidence: str) -> str | None:
    patterns = (
        # This deliberately does not include "responsable del proyecto":
        # a project owner is context, not a task assignee.
        r"(?i)\bresponsable\s*:\s*([A-ZÁÉÍÓÚÜÑ][\wÁÉÍÓÚÜÑáéíóúüñ.'-]+(?:\s+[A-ZÁÉÍÓÚÜÑ][\wÁÉÍÓÚÜÑáéíóúüñ.'-]+){0,3})",
        r"(?i)\b([A-ZÁÉÍÓÚÜÑ][\wÁÉÍÓÚÜÑáéíóúüñ.'-]+(?:\s+[A-ZÁÉÍÓÚÜÑ][\wÁÉÍÓÚÜÑáéíóúüñ.'-]+){0,3})\s+(?:deberá|debera|deber[ií]a|se\s+compromete|se\s+encarga\s+de|confirmará|confirmara|coordinará|coordinara|realizará|realizara|ejecutará|ejecutara|va\s+a|tiene\s+que)\b",
        r"(?i)\b(?:asignad[oa]\s+a|a\s+cargo\s+de)\s+([A-ZÁÉÍÓÚÜÑ][\wÁÉÍÓÚÜÑáéíóúüñ.'-]+(?:\s+[A-ZÁÉÍÓÚÜÑ][\wÁÉÍÓÚÜÑáéíóúüñ.'-]+){0,3})\b",
    )
    for pattern in patterns:
        match = re.search(pattern, evidence)
        if match:
            return _clean(match.group(1), 200)
    return None


def _date_from_evidence(evidence: str) -> str | None:
    for pattern in (r"\b(20\d{2})-(\d{1,2})-(\d{1,2})\b", r"\b(\d{1,2})/(\d{1,2})/(20\d{2})\b"):
        match = re.search(pattern, evidence)
        if not match:
            continue
        try:
            if pattern.startswith(r"\b(20"):
                value = date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
            else:
                value = date(int(match.group(3)), int(match.group(2)), int(match.group(1)))
            return value.isoformat()
        except ValueError:
            pass
    months = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
              "julio": 7, "agosto": 8, "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}
    match = re.search(r"(?i)\b(\d{1,2})\s+de\s+(" + "|".join(months) + r")\s+de\s+(20\d{2})\b", evidence)
    if match:
        try:
            return date(int(match.group(3)), months[match.group(2).casefold()], int(match.group(1))).isoformat()
        except ValueError:
            return None
    return None


def _tracking_from_evidence(evidence: str, options: list[dict[str, str]]) -> tuple[str, str | None]:
    by_name = {str(option.get("label") or option.get("value") or "").casefold(): option for option in options}
    explicit = re.search(r"(?i)tipo\s+de\s+seguimiento\s*:\s*([^.;\n]+)", evidence)
    candidates = [explicit.group(1).strip()] if explicit else []
    candidates.extend(label for label in by_name if label and label in evidence.casefold())
    for candidate in candidates:
        option = by_name.get(candidate.casefold())
        if option:
            label = str(option.get("label") or option.get("value"))
            return label, option.get("jira_id")
    return "Por definir", None


_AREA_CUES = re.compile(
    r"^(?:(?i:el|la)\s+)?(?:(?i:área|area|equipo|sistemas|soporte|departamento|dirección|direccion|unidad|"
    r"operaciones|infraestructura|desarrollo|seguridad)\b|[A-Z]{2,8}(?:-[A-Z]{2,8})?\b)"
)
_AREA_MENTION = re.compile(
    r"\b((?i:área|area|equipo|sistemas|soporte|departamento|dirección|direccion|unidad|"
    r"operaciones|infraestructura|desarrollo|seguridad)"
    r"(?:\s+(?:de\s+)?[A-ZÁÉÍÓÚÜÑ][\wÁÉÍÓÚÜÑáéíóúüñ.-]*){0,3}|"
    r"[A-Z]{2,8}(?:-[A-Z]{2,8})?)\b"
)


def _area_from_evidence(evidence: str) -> dict[str, Any]:
    """Only assign an area when the source explicitly attributes ownership."""
    raw_mentions: list[str] = []
    for match in _AREA_MENTION.finditer(evidence):
        value = _clean(match.group(1), 120)
        if value and value.casefold() not in {v.casefold() for v in raw_mentions}:
            raw_mentions.append(value)

    explicit = re.search(
        r"(?i)\b(?:área\s+responsable|area\s+responsable|equipo\s+responsable|"
        r"responsable\s+del\s+área|responsable\s+del\s+area)\s*:?\s*([^.;\n]+)", evidence,
    )
    attributed = explicit or re.search(
        r"(?i)\b(?:a\s+cargo\s+de|por\s+parte\s+de|por|liderad[oa]s?\s+por|"
        r"ejecutad[oa]s?\s+por|realizad[oa]s?\s+por)\s+([^.;\n]+)", evidence,
    )
    area = _clean(attributed.group(1), 120) if attributed else None
    # A bare personal name after a generic attribution is not an area. An
    # explicitly labeled area remains valid even when its name is unfamiliar.
    if area and not explicit and not _AREA_CUES.search(area):
        area = None
    candidates = [part.strip() for part in re.split(r"\s*(?:,|;|\by\b|\be\b)\s*", area or "") if part.strip()]
    if explicit and len(candidates) == 1:
        status = "EXPLICIT"
        assigned = candidates[0]
    elif len(candidates) > 1:
        status, assigned = "AMBIGUOUS", None
    elif area:
        status, assigned = "EXPLICIT", area
        candidates = [area]
    else:
        status, assigned = ("AMBIGUOUS", None) if raw_mentions else ("UNSPECIFIED", None)
    return {
        "area_responsable": assigned,
        "area_source": "LOCAL",
        "area_status": status,
        "area_candidates": candidates or raw_mentions,
        "area_mentions": raw_mentions,
    }


def _validate_row(row: dict[str, Any], tracking_options: list[dict[str, str]]) -> None:
    if row.get("creation_status") == "CREATED_VERIFIED":
        row["field_status"] = row.get("field_status") or {}
        row["validation_errors"] = []
        row["row_errors"] = []
        row["row_warnings"] = []
        row["row_status"] = "CREATED"
        return
    if row.get("creation_status") in {"CREATED_UNVERIFIED", "UNKNOWN", "REJECTED"}:
        row["row_status"] = "CREATE_ERROR"
        return
    allowed = {str(option.get("label") or option.get("value") or "").casefold() for option in tracking_options}
    statuses: dict[str, str] = {}
    errors: list[dict[str, str]] = []
    required = (("actividad", "Captura una actividad."), ("responsable", "Responsable por definir."),
                ("tipo_de_seguimiento", "Tipo de seguimiento no seleccionado."))
    for field, message in required:
        value = str(row.get(field) or "").strip()
        if not value or value.casefold() == "por definir":
            statuses[field] = "REVIEW"
            errors.append({"field": field, "code": "REQUIRED_REVIEW", "message": message})
        else:
            statuses[field] = "VALID"
    if row.get("responsable") and row.get("responsable") != "Por definir" and row.get("responsable_source") == "MANUAL":
        statuses["responsable"] = "REVIEW"
        errors.append({"field": "responsable", "code": "ASSIGNEE_UNVERIFIED",
                       "message": "Selecciona una coincidencia Jira o conserva un responsable explícito de la minuta."})
    if row.get("responsable_account_id"):
        statuses["responsable"] = "VALID"
    elif row.get("responsable_resolution_status") == "MULTIPLE_MATCHES":
        statuses["responsable"] = "REVIEW"
        errors.append({"field": "responsable", "code": "ASSIGNEE_AMBIGUOUS",
                       "message": "Hay varias coincidencias Jira; selecciona una."})
    tracking = str(row.get("tipo_de_seguimiento") or "").strip()
    if tracking and tracking.casefold() != "por definir" and allowed and tracking.casefold() not in allowed:
        statuses["tipo_de_seguimiento"] = "INVALID"
        errors.append({"field": "tipo_de_seguimiento", "code": "UNKNOWN_TRACKING_TYPE",
                       "message": "El valor no pertenece al catálogo Jira disponible."})
    elif tracking and tracking.casefold() != "por definir" and not allowed:
        statuses["tipo_de_seguimiento"] = "REVIEW"
        errors.append({"field": "tipo_de_seguimiento", "code": "TRACKING_CATALOG_UNAVAILABLE",
                       "message": "No se pudo validar contra el catálogo Jira disponible."})
    elif tracking and tracking.casefold() != "por definir":
        option = next((item for item in tracking_options
                       if str(item.get("label") or item.get("value") or "").casefold() == tracking.casefold()), None)
        if option and not row.get("tipo_de_seguimiento_id"):
            row["tipo_de_seguimiento_id"] = option.get("jira_id")
    raw_date = row.get("fecha_objetivo")
    if raw_date:
        try:
            date.fromisoformat(str(raw_date))
            statuses["fecha_objetivo"] = "VALID"
        except ValueError:
            statuses["fecha_objetivo"] = "INVALID"
            errors.append({"field": "fecha_objetivo", "code": "INVALID_DATE", "message": "Usa una fecha válida."})
    else:
        statuses["fecha_objetivo"] = "EMPTY_ALLOWED"
    for field in ("proyecto_o_frente", "contexto_o_resultado", "comentarios"):
        statuses[field] = "VALID" if row.get(field) else "EMPTY_ALLOWED"
    statuses["area_responsable"] = "VALID" if row.get("area_responsable") else "EMPTY_ALLOWED"
    row["field_status"] = statuses
    row["validation_errors"] = errors
    row["row_errors"] = copy.deepcopy(errors)
    row["row_warnings"] = (
        ["Hay varias áreas mencionadas; confirma manualmente cuál es responsable."]
        if row.get("area_status") == "AMBIGUOUS" else []
    )
    if any(value == "INVALID" for value in statuses.values()):
        row["row_status"] = "INVALID"
    elif any(value == "REVIEW" for value in statuses.values()):
        row["row_status"] = "NEEDS_REVIEW"
    else:
        row["row_status"] = "READY"


def _revalidate(draft: dict[str, Any]) -> dict[str, Any]:
    for row in draft["rows"]:
        _validate_row(row, draft.get("tracking_options", []))
    invalid = sum(any(v == "INVALID" for v in row["field_status"].values()) for row in draft["rows"])
    review = sum(any(v == "REVIEW" for v in row["field_status"].values()) for row in draft["rows"])
    complete = sum(row.get("row_status") == "READY" for row in draft["rows"])
    created = sum(row.get("row_status") == "CREATED" for row in draft["rows"])
    draft["validation_summary"] = {
        "activities": len(draft["rows"]), "complete": complete, "created": created,
        "needs_review": review, "technical_errors": invalid,
    }
    draft["state"] = "INVALID" if invalid else "NEEDS_REVIEW" if review else "READY"
    return draft


class MinuteDraftRepository:
    def __init__(self, database_path: Path):
        self.database_path = database_path
        with sqlite3.connect(database_path) as connection:
            connection.execute("""CREATE TABLE IF NOT EXISTS minute_excel_drafts (
                draft_id TEXT PRIMARY KEY, owner TEXT NOT NULL, version INTEGER NOT NULL,
                created_at TEXT NOT NULL, updated_at TEXT NOT NULL, document TEXT NOT NULL
            )""")
            connection.execute("CREATE INDEX IF NOT EXISTS idx_minute_excel_drafts_owner ON minute_excel_drafts(owner, updated_at DESC)")

    def save_new(self, draft: dict[str, Any]) -> dict[str, Any]:
        with sqlite3.connect(self.database_path) as connection:
            connection.execute("INSERT INTO minute_excel_drafts VALUES (?, ?, ?, ?, ?, ?)",
                               (draft["draft_id"], draft["owner"], draft["version"], draft["created_at"],
                                draft["updated_at"], json.dumps(draft, ensure_ascii=False)))
        return draft

    def get(self, draft_id: str, owner: str) -> dict[str, Any] | None:
        with sqlite3.connect(self.database_path) as connection:
            row = connection.execute("SELECT document FROM minute_excel_drafts WHERE draft_id=? AND owner=?",
                                     (draft_id, owner)).fetchone()
        return json.loads(row[0]) if row else None

    def latest(self, owner: str) -> dict[str, Any] | None:
        with sqlite3.connect(self.database_path) as connection:
            row = connection.execute("SELECT document FROM minute_excel_drafts WHERE owner=? ORDER BY updated_at DESC LIMIT 1",
                                     (owner,)).fetchone()
        return json.loads(row[0]) if row else None

    def replace(self, draft: dict[str, Any], expected_version: int) -> dict[str, Any]:
        with sqlite3.connect(self.database_path) as connection:
            cursor = connection.execute("""UPDATE minute_excel_drafts SET version=?, updated_at=?, document=?
                WHERE draft_id=? AND owner=? AND version=?""",
                (draft["version"], draft["updated_at"], json.dumps(draft, ensure_ascii=False),
                 draft["draft_id"], draft["owner"], expected_version))
            if cursor.rowcount != 1:
                raise MinuteDraftError("El borrador cambió en otra edición. Recarga antes de continuar.", 409, "STALE_DRAFT")
        return draft


class ResponsibleAreaMapRepository:
    """Explicit local source of truth for accountId -> responsible areas.

    Jira Spaces, issue history and project roles are intentionally not used as
    area membership. Operators maintain this small persistent catalogue with
    verified account IDs obtained from Jira's read-only assignee lookup.
    """

    def __init__(self, database_path: Path):
        self.database_path = database_path
        with sqlite3.connect(database_path) as connection:
            connection.execute("""CREATE TABLE IF NOT EXISTS responsible_area_map (
                account_id TEXT NOT NULL,
                area_id TEXT NOT NULL,
                area_display_name TEXT NOT NULL,
                active INTEGER NOT NULL DEFAULT 1,
                source TEXT NOT NULL DEFAULT 'LOCAL_CATALOG',
                updated_at TEXT NOT NULL,
                PRIMARY KEY (account_id, area_id)
            )""")
            connection.execute("CREATE INDEX IF NOT EXISTS idx_responsible_area_map_account ON responsible_area_map(account_id, active)")

    def areas_for(self, account_id: str) -> list[dict[str, str]]:
        with sqlite3.connect(self.database_path) as connection:
            rows = connection.execute("""SELECT area_id, area_display_name, source FROM responsible_area_map
                WHERE account_id=? AND active=1 ORDER BY area_display_name COLLATE NOCASE""", (account_id,)).fetchall()
        return [{"area_id": str(row[0]), "label": str(row[1]), "source": str(row[2])} for row in rows]

    def replace_account(self, account_id: str, areas: list[dict[str, str]], *, source: str = "LOCAL_CATALOG") -> None:
        """Administrative/test maintenance API; never contacts or writes to Jira."""
        now = _now()
        normalized = []
        for area in areas:
            area_id = _clean(area.get("area_id") or area.get("id"), 120)
            label = _clean(area.get("label") or area.get("area_display_name"), 120)
            if area_id and label:
                normalized.append((area_id, label))
        with sqlite3.connect(self.database_path) as connection:
            connection.execute("DELETE FROM responsible_area_map WHERE account_id=?", (account_id,))
            connection.executemany("""INSERT INTO responsible_area_map
                (account_id, area_id, area_display_name, active, source, updated_at) VALUES (?, ?, ?, 1, ?, ?)""",
                [(account_id, area_id, label, source, now) for area_id, label in normalized])

class MinuteDraftService:
    def __init__(self, database_path: Path, proposal_service, template_path: Path):
        self.repository = MinuteDraftRepository(database_path)
        self.area_mappings = ResponsibleAreaMapRepository(database_path)
        self._area_cache: dict[str, list[dict[str, str]]] = {}
        self.proposals = proposal_service
        self.template_path = template_path

    def _areas_for_account(self, account_id: str) -> list[dict[str, str]]:
        if account_id not in self._area_cache:
            self._area_cache[account_id] = self.area_mappings.areas_for(account_id)
        return copy.deepcopy(self._area_cache[account_id])

    def replace_responsible_areas(self, account_id: str, areas: list[dict[str, str]]) -> None:
        """Maintain local mappings and invalidate only that account's cache."""
        if not _clean(account_id, 200):
            raise MinuteDraftError("La cuenta Jira es obligatoria.", 422, "AREA_MAPPING_ACCOUNT_REQUIRED")
        self.area_mappings.replace_account(account_id, areas)
        self._area_cache.pop(account_id, None)

    def apply_responsible_area_resolution(self, draft_id: str, candidate_id: str, owner: str,
                                          expected_version: int) -> dict[str, Any]:
        """Resolve the local area mapping only after a verified Jira account exists."""
        draft = copy.deepcopy(self.get(draft_id, owner))
        if draft["version"] != expected_version:
            raise MinuteDraftError("El borrador cambió en otra edición. Recarga antes de continuar.", 409, "STALE_DRAFT")
        row = next((item for item in draft["rows"] if item["candidate_id"] == candidate_id), None)
        if row is None:
            raise MinuteDraftError("Actividad no encontrada.", 404, "ROW_NOT_FOUND")
        account_id = str(row.get("responsable_account_id") or "").strip()
        if not account_id or row.get("responsable_resolution_status") not in {"AUTO_RESOLVED", "VALIDATED_JIRA"}:
            row.update(area_responsable=None, area_mapping_id=None, area_candidates=[],
                       area_status="UNSPECIFIED", area_source="UNRESOLVED")
            return self._commit(draft, expected_version)
        areas = self._areas_for_account(account_id)
        labels = [item["label"] for item in areas]
        previous = str(row.get("area_responsable") or "")
        previous_source = row.get("area_source")
        supported = next((item for item in areas if canonical(item["label"]) == canonical(previous)), None)
        evidence = canonical(str(row.get("source_evidence") or ""))
        evidenced = next((item for item in areas if canonical(item["label"]) in evidence), None)
        row["area_candidates"] = labels
        if not areas:
            row.update(area_responsable=None, area_mapping_id=None, area_status="UNSPECIFIED", area_source="UNRESOLVED")
        elif previous_source == "USER_SELECTED" and supported:
            row.update(area_mapping_id=supported["area_id"], area_status="EXPLICIT", area_source="USER_SELECTED")
        elif len(areas) == 1:
            row.update(area_responsable=areas[0]["label"], area_mapping_id=areas[0]["area_id"],
                       area_status="EXPLICIT", area_source="SINGLE_MAPPING")
        elif evidenced:
            row.update(area_responsable=evidenced["label"], area_mapping_id=evidenced["area_id"],
                       area_status="EXPLICIT", area_source="TEXT_CONFIRMED")
        else:
            row.update(area_responsable=None, area_mapping_id=None, area_status="AMBIGUOUS", area_source="UNRESOLVED")
        return self._commit(draft, expected_version)

    def create(self, text: str, owner: str, tracking_options: list[dict[str, str]] | None = None,
               enrichment_enabled: bool = False) -> dict[str, Any]:
        normalized, removed = normalize_text(text)
        if len(normalized) < 40:
            raise MinuteDraftError("Pega al menos 40 caracteres con acuerdos o pendientes.")
        extraction = extract_explicit(normalized)
        activities, _, _ = build_activities(extraction.candidates, None)
        project = _project_from_text(normalized)
        project_responsible, project_responsible_evidence = _project_responsible_from_text(normalized)
        rows = []
        options = tracking_options or []
        for candidate, activity in zip(extraction.candidates, activities):
            responsible = _responsible_from_evidence(candidate.source_evidence)
            tracking, tracking_id = _tracking_from_evidence(candidate.source_evidence, options)
            area = _area_from_evidence(candidate.source_evidence)
            rows.append({
                "candidate_id": candidate.candidate_id, "actividad": activity.title,
                **area,
                "responsable": responsible or "Por definir", "responsable_account_id": None,
                "responsable_source": "EXPLICIT" if responsible else "UNRESOLVED",
                "responsable_resolution_status": "EXPLICIT_MINUTES" if responsible else "UNRESOLVED",
                "tipo_de_seguimiento": tracking, "tipo_de_seguimiento_id": tracking_id,
                "fecha_objetivo": _date_from_evidence(candidate.source_evidence),
                "proyecto_o_frente": project, "contexto_o_resultado": activity.description,
                "comentarios": None, "source_section": candidate.source_section,
                "source_evidence": candidate.source_evidence, "extraction_mode": extraction.extraction_mode,
                "touched_fields": [], "suggestions": [], "ignored_suggestions": [],
            })
        # The UI requests ``output=draft`` directly, so it needs the same safe
        # natural-language fallback as the proposal API. Collaborators stay in
        # local context; they are never converted into Jira custom fields.
        semantic = extract_rule_candidates(normalized) if not rows else []
        for index, candidate in enumerate(semantic, 1):
            collaborators = list(dict.fromkeys(candidate.collaborators))
            context = candidate.context
            if collaborators and collaborators[0] not in context:
                context = (context + " " if context else "") + "Coordinar con " + ", ".join(collaborators) + "."
            area = _area_from_evidence(candidate.evidence)
            if candidate.area:
                # Text evidence can suggest an area, but only the existing
                # Jira account→area mapping can confirm it later.
                area = {"area_responsable": candidate.area, "area_source": "TEXT_EVIDENCE",
                        "area_status": "EXPLICIT", "area_candidates": [candidate.area],
                        "area_mentions": [candidate.area]}
            rows.append({
                "candidate_id": f"SEM-{index:03d}", "actividad": candidate.title,
                **area,
                "responsable": candidate.responsible_name or "Por definir", "responsable_account_id": None,
                "responsable_source": "EXPLICIT" if candidate.responsible_name else "UNRESOLVED",
                "responsable_resolution_status": "EXPLICIT_MINUTES" if candidate.responsible_name else "UNRESOLVED",
                "tipo_de_seguimiento": "Por definir", "tipo_de_seguimiento_id": None,
                "fecha_objetivo": candidate.due_date.isoformat() if candidate.due_date else None,
                "proyecto_o_frente": candidate.project or project, "contexto_o_resultado": context or candidate.evidence,
                "comentarios": None, "source_section": None, "source_evidence": candidate.evidence,
                "extraction_mode": "HYBRID_SEMANTIC", "extraction_source": candidate.source,
                "collaborators": collaborators, "touched_fields": [], "suggestions": [], "ignored_suggestions": [],
                "semantic_evidence": {"activity": candidate.activity_evidence, "area": candidate.area_evidence,
                                      "responsible": candidate.responsible_evidence,
                                      "collaborators": candidate.collaborators_evidence,
                                      "responsible_role": candidate.responsible_role},
            })
        now = _now()
        draft = {
            "draft_id": str(uuid.uuid4()), "owner": owner, "source": "MINUTA", "source_text": normalized,
            "source_hash": hashlib.sha256(normalized.encode()).hexdigest(), "created_at": now, "updated_at": now,
            "version": 1, "rows": rows, "tracking_options": options,
            "jira_space_key": "AT", "catalog_status": "LOADING",
            "document": {"project": project, "date": _date_from_text(normalized).isoformat() if _date_from_text(normalized) else None,
                         "participants": [p.model_dump() for p in _participants_from_text(normalized)],
                         "project_responsible": project_responsible,
                         "project_responsible_evidence": project_responsible_evidence,
                         "project_responsible_account_id": None,
                         "project_responsible_source": "EXPLICIT" if project_responsible else "UNRESOLVED",
                         "project_responsible_resolution_status": "EXPLICIT_MINUTES" if project_responsible else "UNRESOLVED",
                         "project_responsible_lookup_query": None},
            "extraction_mode": "HYBRID_SEMANTIC" if semantic else extraction.extraction_mode, "duplicates_removed": removed,
            "warnings": [*extraction.warnings, *([f"Se omitieron {removed} bloques de texto repetidos."] if removed else [])],
            "enrichment": {"state": "PENDING" if enrichment_enabled and extraction.candidates else "NONE", "base_version": 1,
                           "suggestions": [], "message": "La mejora opcional con IA está en curso." if enrichment_enabled and extraction.candidates else None},
            "proposal_id": None,
        }
        return self.repository.save_new(_revalidate(draft))

    def set_tracking_catalog(self, draft_id: str, owner: str, *, options: list[dict[str, Any]],
                             status: str, field_id: str | None = None) -> dict[str, Any]:
        """Merge live Jira options into untouched draft fields without changing a user's version."""
        with sqlite3.connect(self.repository.database_path) as connection:
            stored = connection.execute("SELECT document FROM minute_excel_drafts WHERE draft_id=? AND owner=?",
                                        (draft_id, owner)).fetchone()
            if stored is None:
                raise MinuteDraftError("Borrador no encontrado.", 404, "DRAFT_NOT_FOUND")
            draft = json.loads(stored[0])
            if status != "READY" and draft.get("tracking_options"):
                status = "STALE"
                options = draft["tracking_options"]
            normalized_options = []
            for option in options:
                label = str(option.get("label") or option.get("value") or "").strip()
                if label:
                    normalized_options.append({"value": label, "label": label,
                                               "jira_id": str(option["jira_id"]) if option.get("jira_id") is not None else None})
            draft["tracking_options"] = normalized_options
            draft["catalog_status"] = status
            draft["tracking_field_id"] = field_id
            if status == "READY":
                for row in draft["rows"]:
                    by_id = "tipo_de_seguimiento" in row.get("touched_fields", [])
                    existing = str(row.get("tipo_de_seguimiento") or "").strip()
                    if by_id and existing:
                        match = next((o for o in normalized_options if o["label"].casefold() == existing.casefold()), None)
                        row["tipo_de_seguimiento_id"] = match.get("jira_id") if match else None
                    elif not by_id:
                        name, option_id = _tracking_from_evidence(row.get("source_evidence", ""), normalized_options)
                        row["tipo_de_seguimiento"] = name
                        row["tipo_de_seguimiento_id"] = option_id
            _revalidate(draft)
            connection.execute("UPDATE minute_excel_drafts SET document=? WHERE draft_id=? AND owner=?",
                               (json.dumps(draft, ensure_ascii=False), draft_id, owner))
        return draft

    def save_enrichment(self, draft_id: str, owner: str, base_version: int, *, state: str,
                        suggestions: list[dict[str, Any]], message: str | None,
                        failure_code: str | None = None, duration_ms: float | None = None) -> dict[str, Any] | None:
        """Store suggestions only when the source version is still current."""
        with sqlite3.connect(self.repository.database_path) as connection:
            row = connection.execute("SELECT version, document FROM minute_excel_drafts WHERE draft_id=? AND owner=?",
                                     (draft_id, owner)).fetchone()
            if row is None:
                return None
            current = json.loads(row[1])
            if row[0] != base_version:
                state, suggestions = "STALE", []
                message, failure_code = "La sugerencia quedó obsoleta por una edición humana.", "MINUTES_ENRICHMENT_STALE"
            by_id: dict[str, list[dict[str, Any]]] = {}
            for suggestion in suggestions:
                by_id.setdefault(suggestion["candidate_id"], []).extend(suggestion.get("fields", []))
            for item in current["rows"]:
                item["suggestions"] = by_id.get(item["candidate_id"], []) if state == "READY" else []
            current["enrichment"] = {"state": state, "base_version": base_version,
                                     "suggestions": suggestions, "message": message,
                                     "failure_code": failure_code, "duration_ms": duration_ms}
            connection.execute("UPDATE minute_excel_drafts SET document=? WHERE draft_id=? AND owner=?",
                               (json.dumps(current, ensure_ascii=False), draft_id, owner))
        return current

    def get(self, draft_id: str, owner: str) -> dict[str, Any]:
        draft = self.repository.get(draft_id, owner)
        if draft is None:
            raise MinuteDraftError("Borrador no encontrado.", 404, "DRAFT_NOT_FOUND")
        self._upgrade_local_fields(draft)
        return _revalidate(draft)

    def latest(self, owner: str) -> dict[str, Any] | None:
        draft = self.repository.latest(owner)
        if draft is None:
            return None
        self._upgrade_local_fields(draft)
        return _revalidate(draft)

    @staticmethod
    def _upgrade_local_fields(draft: dict[str, Any]) -> None:
        document = draft.setdefault("document", {})
        document.setdefault("project_responsible", None)
        document.setdefault("project_responsible_evidence", None)
        document.setdefault("project_responsible_account_id", None)
        document.setdefault("project_responsible_source", "UNRESOLVED")
        document.setdefault("project_responsible_resolution_status", "UNRESOLVED")
        document.setdefault("project_responsible_lookup_query", None)
        for row in draft.get("rows", []):
            row.setdefault("area_responsable", None)
            row.setdefault("area_source", "UNRESOLVED")
            row.setdefault("area_status", "UNSPECIFIED")
            row.setdefault("area_candidates", [])
            row.setdefault("area_mentions", [])
            row.setdefault("area_mapping_id", None)
            row.setdefault("proposal_id", None)
            row.setdefault("creation_status", None)

    def _commit(self, draft: dict[str, Any], expected_version: int) -> dict[str, Any]:
        draft["version"] = expected_version + 1
        draft["updated_at"] = _now()
        draft["proposal_id"] = None
        if draft.get("enrichment", {}).get("state") == "PENDING":
            draft["enrichment"] = {**draft["enrichment"], "state": "STALE", "message": "La sugerencia quedó obsoleta por una edición humana."}
        return self.repository.replace(_revalidate(draft), expected_version)

    def update_row(self, draft_id: str, candidate_id: str, owner: str, expected_version: int,
                   changes: dict[str, Any]) -> dict[str, Any]:
        if set(changes) - EDITABLE_FIELDS:
            raise MinuteDraftError("La edición contiene campos no permitidos.")
        if "area_source" in changes and "area_responsable" not in changes:
            raise MinuteDraftError("El origen de área requiere un área seleccionada.")
        if changes.get("area_source") not in {None, "LOCAL", "JIRA", "SINGLE_MAPPING", "TEXT_CONFIRMED", "USER_SELECTED", "MANUAL", "UNRESOLVED"}:
            raise MinuteDraftError("Origen de área inválido.")
        draft = copy.deepcopy(self.get(draft_id, owner))
        if draft["version"] != expected_version:
            raise MinuteDraftError("El borrador cambió en otra edición. Recarga antes de continuar.", 409, "STALE_DRAFT")
        row = next((item for item in draft["rows"] if item["candidate_id"] == candidate_id), None)
        if row is None:
            raise MinuteDraftError("Actividad no encontrada.", 404, "ROW_NOT_FOUND")
        if row.get("creation_status") in {"CREATED_VERIFIED", "CREATED_UNVERIFIED", "UNKNOWN"}:
            raise MinuteDraftError("La actividad ya fue enviada o su resultado es incierto; no se puede editar ni reenviar desde el borrador.",
                                   409, "ROW_ALREADY_SUBMITTED")
        limits = {"actividad": 500, "responsable": 200, "tipo_de_seguimiento": 200,
                  "area_responsable": 120, "proyecto_o_frente": 300, "contexto_o_resultado": 4000, "comentarios": 2000}
        for field, value in changes.items():
            row[field] = _clean(value, limits[field]) if field in limits else (value if field == "ignored_suggestions" else value or None)
            if field not in row["touched_fields"]:
                row["touched_fields"].append(field)
        if "area_responsable" in changes:
            row["area_source"] = changes.get("area_source") or "MANUAL"
            row["area_status"] = "EXPLICIT" if row.get("area_responsable") else ("AMBIGUOUS" if row.get("area_candidates") else "UNSPECIFIED")
            if row.get("area_responsable"):
                matched = next((area for area in self._areas_for_account(str(row.get("responsable_account_id") or ""))
                                if canonical(str(area.get("label") or "")) == canonical(str(row["area_responsable"]))), None)
                row["area_mapping_id"] = matched.get("area_id") if matched else None
        row["proposal_id"] = None
        row.pop("proposal_status", None)
        resolution_status = changes.get("responsable_resolution_status")
        if resolution_status and resolution_status not in RESPONSABLE_STATES:
            raise MinuteDraftError("Estado de responsable inválido.")
        if "tipo_de_seguimiento" in changes:
            tracking = str(row.get("tipo_de_seguimiento") or "").strip()
            option = next((item for item in draft.get("tracking_options", [])
                           if str(item.get("label") or item.get("value") or "").casefold() == tracking.casefold()), None)
            row["tipo_de_seguimiento_id"] = option.get("jira_id") if option else None
        elif "tipo_de_seguimiento_id" in changes:
            option = next((item for item in draft.get("tracking_options", [])
                           if item.get("jira_id") == row.get("tipo_de_seguimiento_id")
                           and str(item.get("label") or item.get("value") or "").casefold() == str(row.get("tipo_de_seguimiento") or "").casefold()), None)
            if not option:
                raise MinuteDraftError("El tipo de seguimiento no coincide con una opción Jira permitida.", 422, "TRACKING_OPTION_NOT_VERIFIED")
        if "responsable" in changes:
            if changes.get("responsable_account_id"):
                row["responsable_source"] = "JIRA"
                row["responsable_resolution_status"] = resolution_status or "VALIDATED_JIRA"
            elif str(row.get("responsable") or "").casefold() == "por definir":
                row["responsable_source"] = "UNRESOLVED"
                row["responsable_resolution_status"] = "UNRESOLVED"
                row["responsable_account_id"] = None
            else:
                if resolution_status in {"AUTO_RESOLVED", "VALIDATED_JIRA"}:
                    raise MinuteDraftError("Una validación Jira requiere una cuenta verificada.", 422, "ASSIGNEE_NOT_VERIFIED")
                extracted = _responsible_from_evidence(row.get("source_evidence", ""))
                is_explicit = bool(extracted and canonical(extracted) == canonical(str(row.get("responsable") or "")))
                row["responsable_source"] = "EXPLICIT" if is_explicit else "MANUAL"
                row["responsable_resolution_status"] = (
                    resolution_status or ("EXPLICIT_MINUTES" if is_explicit else "NEEDS_REVIEW")
                )
                row["responsable_account_id"] = None
        elif resolution_status:
            if resolution_status in {"AUTO_RESOLVED", "VALIDATED_JIRA"} and not row.get("responsable_account_id"):
                raise MinuteDraftError("Una validación Jira requiere una cuenta verificada.", 422, "ASSIGNEE_NOT_VERIFIED")
            row["responsable_resolution_status"] = resolution_status
        return self._commit(draft, expected_version)

    def update_project_responsible(self, draft_id: str, owner: str, expected_version: int,
                                   changes: dict[str, Any]) -> dict[str, Any]:
        """Edit document ownership without changing any activity assignee."""
        if set(changes) - PROJECT_RESPONSIBLE_EDITABLE_FIELDS:
            raise MinuteDraftError("La edición del responsable del proyecto contiene campos no permitidos.")
        draft = copy.deepcopy(self.get(draft_id, owner))
        if draft["version"] != expected_version:
            raise MinuteDraftError("El borrador cambió en otra edición. Recarga antes de continuar.", 409, "STALE_DRAFT")
        document = draft.setdefault("document", {})
        for field, value in changes.items():
            document[field] = _clean(value, 200) if field == "project_responsible" else (value or None)
        status = document.get("project_responsible_resolution_status")
        if status and status not in RESPONSABLE_STATES:
            raise MinuteDraftError("Estado del responsable del proyecto inválido.")
        if document.get("project_responsible_account_id"):
            document["project_responsible_source"] = "JIRA"
            document["project_responsible_resolution_status"] = status or "VALIDATED_JIRA"
        elif not document.get("project_responsible"):
            document.update(project_responsible_source="UNRESOLVED", project_responsible_resolution_status="UNRESOLVED")
        else:
            if status in {"AUTO_RESOLVED", "VALIDATED_JIRA"}:
                raise MinuteDraftError("Una validación Jira requiere una cuenta verificada.", 422, "ASSIGNEE_NOT_VERIFIED")
            document["project_responsible_account_id"] = None
            document["project_responsible_source"] = "EXPLICIT" if document.get("project_responsible_evidence") else "MANUAL"
            document["project_responsible_resolution_status"] = status or "EXPLICIT_MINUTES"
        return self._commit(draft, expected_version)

    def add_row(self, draft_id: str, owner: str, expected_version: int) -> dict[str, Any]:
        draft = copy.deepcopy(self.get(draft_id, owner))
        if draft["version"] != expected_version:
            raise MinuteDraftError("El borrador cambió en otra edición. Recarga antes de continuar.", 409, "STALE_DRAFT")
        ids = {row["candidate_id"] for row in draft["rows"]}
        number = 1
        while f"MAN-{number:03d}" in ids:
            number += 1
        draft["rows"].append({"candidate_id": f"MAN-{number:03d}", "actividad": None,
            "responsable": "Por definir", "responsable_account_id": None, "tipo_de_seguimiento": "Por definir",
            "responsable_source": "UNRESOLVED", "responsable_resolution_status": "UNRESOLVED",
            "tipo_de_seguimiento_id": None, "fecha_objetivo": None, "proyecto_o_frente": draft["document"].get("project"),
            "contexto_o_resultado": None, "comentarios": None, "source_section": "Fila manual",
            "source_evidence": "Actividad agregada manualmente por el usuario.", "extraction_mode": "MANUAL",
            "touched_fields": [], "suggestions": [], "ignored_suggestions": []})
        return self._commit(draft, expected_version)

    def delete_row(self, draft_id: str, candidate_id: str, owner: str, expected_version: int) -> dict[str, Any]:
        draft = copy.deepcopy(self.get(draft_id, owner))
        if draft["version"] != expected_version:
            raise MinuteDraftError("El borrador cambió en otra edición. Recarga antes de continuar.", 409, "STALE_DRAFT")
        before = len(draft["rows"])
        target = next((row for row in draft["rows"] if row["candidate_id"] == candidate_id), None)
        if target and target.get("creation_status") in {"CREATED_VERIFIED", "CREATED_UNVERIFIED", "UNKNOWN"}:
            raise MinuteDraftError("La actividad ya fue enviada o su resultado es incierto y debe conservarse en el borrador.",
                                   409, "ROW_ALREADY_SUBMITTED")
        draft["rows"] = [row for row in draft["rows"] if row["candidate_id"] != candidate_id]
        if len(draft["rows"]) == before:
            raise MinuteDraftError("Actividad no encontrada.", 404, "ROW_NOT_FOUND")
        return self._commit(draft, expected_version)

    def validate(self, draft_id: str, owner: str, tracking_options: list[dict[str, str]] | None = None) -> dict[str, Any]:
        draft = copy.deepcopy(self.get(draft_id, owner))
        if tracking_options is not None:
            draft["tracking_options"] = tracking_options
        return _revalidate(draft)

    def export_xlsx(self, draft_id: str, owner: str) -> bytes:
        draft = self.get(draft_id, owner)
        workbook = load_workbook(self.template_path)
        sheet = workbook["PLAN_DE_TRABAJO"]
        for row_number in range(6, 56):
            for column in range(1, len(HEADERS) + 1):
                sheet.cell(row_number, column).value = None
        for row_number, row in enumerate(draft["rows"][:50], 6):
            values: list[Any] = [row.get("actividad"), row.get("area_responsable"), row.get("responsable"),
                                 row.get("tipo_de_seguimiento"),
                                 date.fromisoformat(row["fecha_objetivo"]) if row.get("fecha_objetivo") else None,
                                 row.get("proyecto_o_frente"), row.get("contexto_o_resultado"), row.get("comentarios")]
            for column, value in enumerate(values, 1):
                if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                    value = "'" + value
                sheet.cell(row_number, column).value = value
            if values[4]:
                sheet.cell(row_number, 5).number_format = "yyyy-mm-dd"
        state_sheet = workbook["ESTADO_LOCAL"] if "ESTADO_LOCAL" in workbook.sheetnames else workbook.create_sheet("ESTADO_LOCAL")
        state_sheet.delete_rows(1, state_sheet.max_row)
        state_headers = (
            "Candidate_ID", "Estado", "Errores", "Advertencias", "Área_origen", "Área_estado",
            "Evidencia_origen", "Proposal_ID", "Creation_status", "Jira_issue_key", "Jira_issue_url", "Created_at",
        )
        state_sheet.append(state_headers)
        for row in draft["rows"][:50]:
            values = tuple(
                "'" + value if isinstance(value, str) and (
                    value.lstrip().startswith(("=", "+", "@"))
                    or (value.lstrip().startswith("-") and len(value.lstrip()) > 1 and not value.lstrip()[1].isspace())
                ) else value
                for value in (
                row.get("candidate_id"), row.get("row_status"),
                " | ".join(issue["message"] for issue in row.get("row_errors", [])),
                " | ".join(row.get("row_warnings", [])), row.get("area_source"), row.get("area_status"),
                row.get("source_evidence"), row.get("proposal_id"), row.get("creation_status"),
                row.get("jira_issue_key"), row.get("jira_issue_url"), row.get("created_at"),
                )
            )
            state_sheet.append(values)
        state_sheet.freeze_panes = "A2"
        state_sheet.auto_filter.ref = state_sheet.dimensions
        for cell in state_sheet[1]:
            font = copy.copy(cell.font)
            font.bold = True
            cell.font = font
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def to_proposal(self, draft_id: str, owner: str, candidate_ids: list[str] | None = None) -> dict[str, Any]:
        draft = self.validate(draft_id, owner)
        rows_by_id = {row["candidate_id"]: row for row in draft["rows"]}
        requested = candidate_ids if candidate_ids is not None else [
            row["candidate_id"] for row in draft["rows"]
            if row.get("row_status") == "READY" and not row.get("proposal_id")
        ]
        selected = list(dict.fromkeys(requested))
        if not selected:
            raise MinuteDraftError("No hay actividades listas sin propuesta para preparar.", 409, "NO_READY_ROWS")
        if candidate_ids and len(selected) != len(candidate_ids):
            raise MinuteDraftError("La selección contiene actividades repetidas.", 422, "DUPLICATE_CANDIDATE")
        if any(candidate_id not in rows_by_id for candidate_id in selected):
            raise MinuteDraftError("La selección incluye una actividad que no existe.", 404, "ROW_NOT_FOUND")
        selected_rows = [rows_by_id[candidate_id] for candidate_id in selected]
        if all(row.get("proposal_id") for row in selected_rows):
            proposal_ids = {row["proposal_id"] for row in selected_rows}
            if len(proposal_ids) == 1:
                existing = self.proposals.repository.get(next(iter(proposal_ids)))
                if existing:
                    return existing
        if any(row.get("proposal_id") for row in selected_rows):
            raise MinuteDraftError("Quita de la selección las actividades que ya tienen una propuesta.", 409, "ROW_ALREADY_PROPOSED")
        if any(row.get("row_status") != "READY" for row in selected_rows):
            raise MinuteDraftError("Sólo puedes preparar actividades listas. Las demás siguen en revisión.", 409, "ROW_NOT_READY")
        selected_draft = copy.deepcopy(draft)
        selected_draft["rows"] = selected_rows
        for index, row in enumerate(selected_draft["rows"]):
            row["source_row"] = index + 6
        proposal = self.proposals.create_from_minute_draft(selected_draft, owner)
        current = copy.deepcopy(self.get(draft_id, owner))
        expected = current["version"]
        for row in current["rows"]:
            if row["candidate_id"] in selected:
                row["proposal_id"] = proposal["proposal_id"]
                row["proposal_status"] = "PROPOSAL_CREATED"
        current["proposal_id"] = proposal["proposal_id"]
        current["version"] = expected + 1
        current["updated_at"] = _now()
        self.repository.replace(current, expected)
        return proposal

    def sync_proposal_execution(self, proposal_id: str, owner: str) -> dict[str, Any] | None:
        """Reflect protected Objective 4C.1 outcomes back onto their source draft rows."""
        proposal = self.proposals.repository.get(proposal_id)
        if not proposal or proposal.get("created_by") != owner or not proposal.get("minutes"):
            return None
        document = proposal["minutes"].get("document") or {}
        draft_id = document.get("draft_id")
        candidate_ids = document.get("candidate_ids") or []
        if not draft_id or len(candidate_ids) != len(proposal.get("items", [])):
            return None
        row_id_to_candidate = {
            item["generated"]["row_id"]: candidate_id
            for item, candidate_id in zip(proposal["items"], candidate_ids)
        }
        outcomes: dict[str, dict[str, Any]] = {}
        intents = self.proposals.repository.list_execution_intents(proposal_id)
        for intent in reversed(intents):
            if intent.get("objective") not in {"4C.1", "4D", "4E"}:
                continue
            for item in intent.get("items", []):
                candidate_id = row_id_to_candidate.get(item.get("row_id"))
                if candidate_id and item.get("status") in {"CREATED_VERIFIED", "CREATED_UNVERIFIED", "REJECTED", "UNKNOWN"}:
                    outcomes[candidate_id] = {
                        "creation_status": item["status"], "jira_issue_key": item.get("issue_key"),
                        "jira_issue_url": item.get("issue_url"), "created_at": item.get("completed_at"),
                        "creation_idempotency_key": item.get("idempotency_key"),
                    }
        if not outcomes:
            return self.get(draft_id, owner)
        current = self.get(draft_id, owner)
        changed = False
        for row in current["rows"]:
            result = outcomes.get(row["candidate_id"])
            if result and any(row.get(field) != value for field, value in result.items()):
                row.update(result)
                changed = True
        if not changed:
            return current
        expected = current["version"]
        current["version"] = expected + 1
        current["updated_at"] = _now()
        _revalidate(current)
        return self.repository.replace(current, expected)
