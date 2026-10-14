from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .errors import ValidationIssue, WorkbookValidationError


CAPTURE_SHEET = "PLAN_DE_TRABAJO"
HEADER_ROW = 5
FIRST_DATA_ROW = 6
LAST_DATA_ROW = 55
EXPECTED_HEADERS = (
    "Actividad *",
    "Responsable *",
    "Frente_de_trabajo_o_Proyecto",
    "Contexto_o_Resultado",
    "Fecha_Objetivo",
    "Estado_Actual",
    "Comentarios",
)
EXPECTED_HEADERS_V12 = (
    "Actividad *",
    "Responsable *",
    "Tipo_de_seguimiento *",
    "Fecha_Objetivo",
    "Proyecto_o_Frente",
    "Contexto_o_Resultado",
    "Comentarios",
)
EXPECTED_HEADERS_V13 = (
    "Actividad *", "Área_Responsable", "Responsable *", "Tipo_de_seguimiento *",
    "Fecha_Objetivo", "Proyecto_o_Frente", "Contexto_o_Resultado", "Comentarios",
)
ALLOWED_STATES = {
    "Pendiente",
    "En curso",
    "En pausa",
    "Completada",
    "Rechazada",
    "Por confirmar",
}
STATE_ALIASES = {
    "pendiente": "Pendiente",
    "en curso": "En curso",
    "en proceso": "En curso",
    "en pausa": "En pausa",
    "pausada": "En pausa",
    "completada": "Completada",
    "completado": "Completada",
    "terminada": "Completada",
    "terminado": "Completada",
    "rechazada": "Rechazada",
    "rechazado": "Rechazada",
    "cancelada": "Rechazada",
    "cancelado": "Rechazada",
    "por confirmar": "Por confirmar",
}


@dataclass(frozen=True)
class ParsedPlan:
    filename: str
    sha256: str
    template_version: str
    rows: list[dict[str, Any]]


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def normalize_state(value: Any) -> str:
    text = clean_text(value)
    if not text:
        return "Pendiente"
    return STATE_ALIASES.get(text.casefold(), text)


def normalize_date(value: Any, epoch: datetime) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, epoch=epoch)
            return (
                converted.date().isoformat()
                if isinstance(converted, datetime)
                else converted.isoformat()
            )
        except (ValueError, OverflowError, TypeError):
            return None

    text = clean_text(value)
    if not text:
        return None
    for format_string in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(text, format_string).date().isoformat()
        except ValueError:
            continue
    return None


class TemplateParser:
    def parse(self, filename: str, content: bytes) -> ParsedPlan:
        issues: list[ValidationIssue] = []
        if not filename.lower().endswith(".xlsx"):
            raise WorkbookValidationError(
                "El archivo debe tener extensión .xlsx.",
                [ValidationIssue("INVALID_EXTENSION", "Sólo se acepta la plantilla .xlsx.")],
            )
        if not content:
            raise WorkbookValidationError(
                "El archivo está vacío.",
                [ValidationIssue("EMPTY_FILE", "Selecciona un archivo Excel con contenido.")],
            )

        try:
            workbook = load_workbook(io.BytesIO(content), data_only=False, read_only=False)
        except Exception as exc:
            raise WorkbookValidationError(
                "No fue posible abrir el archivo Excel.",
                [ValidationIssue("INVALID_XLSX", f"El archivo no es un .xlsx válido: {exc}")],
            ) from exc

        if CAPTURE_SHEET not in workbook.sheetnames:
            raise WorkbookValidationError(
                "La plantilla no contiene la hoja de captura requerida.",
                [
                    ValidationIssue(
                        "MISSING_SHEET",
                        f"Se esperaba la hoja '{CAPTURE_SHEET}'.",
                    )
                ],
            )

        sheet = workbook[CAPTURE_SHEET]
        version_text = clean_text(sheet["A3"].value) or ""
        version = next((candidate for candidate in ("1.3", "1.2", "1.1")
                        if f"v{candidate}" in version_text.casefold()), "1.1")
        expected_headers = {"1.1": EXPECTED_HEADERS, "1.2": EXPECTED_HEADERS_V12,
                            "1.3": EXPECTED_HEADERS_V13}[version]
        actual_headers = tuple(clean_text(sheet.cell(HEADER_ROW, col).value)
                               for col in range(1, len(expected_headers) + 1))
        if actual_headers != expected_headers:
            for index, (expected, actual) in enumerate(zip(expected_headers, actual_headers), 1):
                if actual != expected:
                    issues.append(
                        ValidationIssue(
                            "INVALID_HEADER",
                            f"Se esperaba '{expected}' y se encontró '{actual or '(vacío)'}'.",
                            row=HEADER_ROW,
                            column=sheet.cell(HEADER_ROW, index).column_letter,
                        )
                    )
            raise WorkbookValidationError(
                f"Los encabezados no corresponden a la Plantilla Excel v{version}.",
                issues,
            )

        if not any(f"v{candidate}" in version_text.casefold() for candidate in ("1.1", "1.2", "1.3")):
            raise WorkbookValidationError(
                        "No se pudo confirmar la versión 1.1, 1.2 o 1.3 de la plantilla.",
                [
                    ValidationIssue(
                        "INVALID_TEMPLATE_VERSION",
                        "La celda A3 debe identificar la plantilla como v1.1, v1.2 o v1.3.",
                        row=3,
                        column="A",
                    )
                ],
            )

        rows: list[dict[str, Any]] = []
        for row_number in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
            values = [sheet.cell(row_number, col).value for col in range(1, len(expected_headers) + 1)]
            activity = clean_text(values[0])
            area_responsable = None
            if version == "1.3":
                area_responsable, responsible, tracking_type = clean_text(values[1]), clean_text(values[2]), clean_text(values[3])
                normalized_date = normalize_date(values[4], workbook.epoch)
                project_or_front, context, comments = clean_text(values[5]), clean_text(values[6]), clean_text(values[7])
                state = "Pendiente"
                optional_values = [area_responsable, tracking_type, values[4], project_or_front, context, comments]
            elif version == "1.2":
                responsible, tracking_type = clean_text(values[1]), clean_text(values[2])
                normalized_date = normalize_date(values[3], workbook.epoch)
                project_or_front, context, comments = clean_text(values[4]), clean_text(values[5]), clean_text(values[6])
                state = "Pendiente"
                optional_values = [tracking_type, values[3], project_or_front, context, comments]
            else:
                responsible = clean_text(values[1])
                tracking_type = None
                normalized_date = normalize_date(values[4], workbook.epoch)
                project_or_front, context, comments = clean_text(values[2]), clean_text(values[3]), clean_text(values[6])
                state = normalize_state(values[5])
                optional_values = [project_or_front, context, values[4], comments]
            if not activity and not responsible and not any(optional_values) and state == "Pendiente":
                continue
            if not activity:
                issues.append(ValidationIssue("MISSING_ACTIVITY", "La actividad es obligatoria.", row=row_number, column="A"))
            if not responsible:
                issues.append(ValidationIssue("MISSING_ASSIGNEE", "El responsable principal es obligatorio.", row=row_number,
                                              column="C" if version == "1.3" else "B"))
            if version in {"1.2", "1.3"} and not tracking_type:
                issues.append(ValidationIssue("MISSING_TRACKING_TYPE", "El tipo de seguimiento es obligatorio.", row=row_number,
                                              column="D" if version == "1.3" else "C"))
            if state not in ALLOWED_STATES:
                issues.append(ValidationIssue("INVALID_STATE", f"El estado '{state}' no pertenece al catálogo v1.1.",
                                              row=row_number, column="F"))
            date_index = 4 if version == "1.3" else 3 if version == "1.2" else 4
            date_value = values[date_index]
            if date_value not in (None, "") and normalized_date is None:
                issues.append(ValidationIssue("INVALID_DATE", "La fecha debe ser una fecha Excel o usar el formato AAAA-MM-DD.",
                                              row=row_number, column="E" if version == "1.3" else "D" if version == "1.2" else "E"))
            if activity and responsible and state in ALLOWED_STATES and (version == "1.1" or tracking_type):
                parsed_row = {
                    "source_row": row_number, "actividad": activity, "responsable_display": responsible,
                    "frente_trabajo_origen": project_or_front, "contexto_resultado": context,
                    "fecha_objetivo": normalized_date, "estado_actual": state, "comentarios": comments,
                    "tipo_seguimiento_origen": tracking_type,
                }
                if version == "1.3":
                    parsed_row.update({"area_responsable": area_responsable, "area_source": "LOCAL",
                                       "area_status": "EXPLICIT" if area_responsable else "UNSPECIFIED",
                                       "area_candidates": [area_responsable] if area_responsable else [],
                                       "area_mentions": []})
                rows.append(parsed_row)

        if issues:
            raise WorkbookValidationError(
                f"Se encontraron {len(issues)} problemas en las filas capturadas.",
                issues,
            )
        if not rows:
            raise WorkbookValidationError(
                "La plantilla no contiene actividades.",
                [
                    ValidationIssue(
                        "NO_ACTIVITIES",
                        "Captura al menos una actividad y un responsable entre las filas 6 y 55.",
                    )
                ],
            )
        return ParsedPlan(
            filename=filename,
            sha256=hashlib.sha256(content).hexdigest(),
            template_version=version,
            rows=rows,
        )
