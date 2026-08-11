from __future__ import annotations

import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import Settings
from app.main import create_app


PROJECT_DIR = Path(__file__).resolve().parents[2]
CANONICAL_TEMPLATE = (
    PROJECT_DIR
    / "releases"
    / "plantilla_excel_v1_1"
    / "plantilla_plan_trabajo_jira_v1_1.xlsx"
)
SCHEMA = PROJECT_DIR / "docs" / "schemas" / "jira_proposal_v1_1.schema.json"


def workbook_bytes(rows: list[list[object]]) -> bytes:
    workbook = load_workbook(CANONICAL_TEMPLATE)
    sheet = workbook["PLAN_DE_TRABAJO"]
    for row_index, values in enumerate(rows, 6):
        for column_index, value in enumerate(values, 1):
            sheet.cell(row_index, column_index).value = value
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


@pytest.fixture
def sample_workbook() -> bytes:
    return workbook_bytes(
        [
            [
                "Implementar filtros para la tabla de solicitudes",
                "Giselle Bello",
                "Mejoras",
                "Permitir localizar solicitudes por estado",
                "2026-08-15",
                "Pendiente",
                None,
            ],
            [
                "BUG: Error de CORS al consultar el servicio",
                "Raúl Gallardo",
                "Integraciones",
                "La petición debe responder sin bloqueo del navegador",
                None,
                "Completada",
                "Corregido previamente en desarrollo",
            ],
            [
                "Definir reglas",
                "Wilfrido / José",
                "Backend",
                None,
                None,
                "Pendiente",
                None,
            ],
        ]
    )


@pytest.fixture
def client(tmp_path: Path):
    settings = Settings(
        database_path=tmp_path / "test.db",
        schema_path=SCHEMA,
        frontend_dist=tmp_path / "dist",
        frozen_template_path=CANONICAL_TEMPLATE,
        local_user="admin@agentejira.local",
        local_password="AgenteJira2026!",
    )
    with TestClient(create_app(settings)) as test_client:
        yield test_client


@pytest.fixture
def auth_headers(client: TestClient) -> dict[str, str]:
    response = client.post(
        "/api/auth/login",
        json={
            "username": "admin@agentejira.local",
            "password": "AgenteJira2026!",
        },
    )
    assert response.status_code == 200
    token = response.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}

