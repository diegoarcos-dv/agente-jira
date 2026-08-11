from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from app.errors import WorkbookValidationError
from app.parser import TemplateParser
from tests.conftest import CANONICAL_TEMPLATE, workbook_bytes


def test_parser_accepts_template_v1_1_and_normalizes_values(sample_workbook: bytes):
    parsed = TemplateParser().parse("plan.xlsx", sample_workbook)

    assert len(parsed.rows) == 3
    assert parsed.rows[0]["source_row"] == 6
    assert parsed.rows[0]["frente_trabajo_origen"] == "Mejoras"
    assert parsed.rows[0]["fecha_objetivo"] == "2026-08-15"
    assert parsed.rows[1]["estado_actual"] == "Completada"
    assert len(parsed.sha256) == 64


def test_parser_rejects_changed_header():
    workbook = load_workbook(CANONICAL_TEMPLATE)
    sheet = workbook["PLAN_DE_TRABAJO"]
    sheet["A5"] = "Título"
    sheet["A6"] = "Actividad de prueba"
    sheet["B6"] = "Giselle Bello"
    stream = io.BytesIO()
    workbook.save(stream)

    with pytest.raises(WorkbookValidationError) as caught:
        TemplateParser().parse("modificada.xlsx", stream.getvalue())

    assert caught.value.issues[0].code == "INVALID_HEADER"
    assert caught.value.issues[0].row == 5
    assert caught.value.issues[0].column == "A"


def test_parser_rejects_missing_capture_sheet():
    workbook = load_workbook(CANONICAL_TEMPLATE)
    del workbook["PLAN_DE_TRABAJO"]
    stream = io.BytesIO()
    workbook.save(stream)

    with pytest.raises(WorkbookValidationError) as caught:
        TemplateParser().parse("sin-hoja.xlsx", stream.getvalue())

    assert caught.value.issues[0].code == "MISSING_SHEET"


def test_parser_rejects_unconfirmed_template_version():
    workbook = load_workbook(CANONICAL_TEMPLATE)
    sheet = workbook["PLAN_DE_TRABAJO"]
    sheet["A3"] = "Plantilla sin versión"
    sheet["A6"] = "Actividad de prueba"
    sheet["B6"] = "Giselle Bello"
    stream = io.BytesIO()
    workbook.save(stream)

    with pytest.raises(WorkbookValidationError) as caught:
        TemplateParser().parse("otra-version.xlsx", stream.getvalue())

    assert caught.value.issues[0].code == "INVALID_TEMPLATE_VERSION"


def test_parser_reports_missing_required_activity():
    content = workbook_bytes(
        [[None, "Giselle Bello", "Mejoras", "Hay contexto sin actividad", None, "Pendiente", None]]
    )

    with pytest.raises(WorkbookValidationError) as caught:
        TemplateParser().parse("incompleta.xlsx", content)

    assert any(issue.code == "MISSING_ACTIVITY" and issue.row == 6 for issue in caught.value.issues)


def test_parser_ignores_reserved_rows_with_default_pending():
    content = workbook_bytes(
        [["Actividad válida", "Giselle Bello", None, None, None, "Pendiente", None]]
    )
    parsed = TemplateParser().parse("plan.xlsx", content)
    assert [row["source_row"] for row in parsed.rows] == [6]
