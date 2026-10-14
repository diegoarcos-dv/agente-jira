"""Build the v1.2 workbook using only backend dependencies.

The previous generator depended on the Codex-only ``@oai/artifact-tool``
package.  This local generator keeps the same workbook contract while making
the dynamic template endpoint reproducible after a clean install.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parents[1]
BASE_TEMPLATE = (
    PROJECT_DIR
    / "releases"
    / "plantilla_excel_v1_2"
    / "plantilla_plan_trabajo_jira_v1_2.xlsx"
)


def build(catalog_path: Path, output_path: Path) -> None:
    document = json.loads(catalog_path.read_text(encoding="utf-8"))
    catalog = document.get("catalog", document)
    options = catalog.get("options") or []

    workbook = load_workbook(BASE_TEMPLATE)
    plan = workbook["PLAN_DE_TRABAJO"]
    catalogs = workbook["CATALOGOS"]

    # Clear only the catalog payload; the frozen workbook supplies the layout.
    for row in catalogs.iter_rows(min_row=1, max_row=max(catalogs.max_row, len(options) + 8), min_col=1, max_col=5):
        for cell in row:
            cell.value = None

    headers = [["display_value", "jira_option_id"]]
    rows = headers + [[option.get("value"), option.get("jira_id")] for option in options]
    for row_index, values in enumerate(rows, start=1):
        for column_index, value in enumerate(values, start=1):
            catalogs.cell(row_index, column_index).value = value

    metadata = [
        ("metadata", "value"),
        ("template_version", "1.2"),
        ("space", catalog.get("space")),
        ("work_type", catalog.get("work_type")),
        ("field_id", catalog.get("field_id")),
        ("catalog_fetched_at", catalog.get("fetched_at")),
        ("catalog_status", catalog.get("status") or "unavailable"),
    ]
    for row_index, values in enumerate(metadata, start=1):
        catalogs.cell(row_index, 4).value = values[0]
        catalogs.cell(row_index, 5).value = values[1]

    for validation in list(plan.data_validations.dataValidation):
        plan.data_validations.dataValidation.remove(validation)
    if options:
        from openpyxl.worksheet.datavalidation import DataValidation

        validation = DataValidation(
            type="list",
            formula1=f"='CATALOGOS'!$A$2:$A${len(options) + 1}",
            allow_blank=True,
        )
        plan.add_data_validation(validation)
        validation.add("C6:C55")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    manifest = {
        "version": "1.2",
        "sheet": "PLAN_DE_TRABAJO",
        "headers": [
            "Actividad *",
            "Responsable *",
            "Tipo_de_seguimiento *",
            "Fecha_Objetivo",
            "Proyecto_o_Frente",
            "Contexto_o_Resultado",
            "Comentarios",
        ],
        "catalog": catalog,
    }
    output_path.with_name("manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    output_path.with_name("README.md").write_text(
        "# Plantilla Excel Jira v1.2\n\n"
        "Incluye captura guiada con catálogo desplegable para "
        "`Tipo_de_seguimiento` y es compatible con el parser v1.1/v1.2. "
        "Esta copia fue generada desde el catálogo verificado disponible en Jira; "
        "si el estado es `unavailable`, no se inventan opciones.\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    if len(sys.argv) not in (2, 3):
        raise SystemExit("Uso: build_template_v12.py CATALOG.json [OUTPUT.xlsx]")
    catalog_file = Path(sys.argv[1]).resolve()
    output_file = (
        Path(sys.argv[2]).resolve()
        if len(sys.argv) == 3
        else PROJECT_DIR / "releases" / "plantilla_excel_v1_2" / "plantilla_plan_trabajo_jira_v1_2.xlsx"
    )
    build(catalog_file, output_file)
