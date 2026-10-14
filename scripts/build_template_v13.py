"""Build the v1.3 workbook with an optional local area metadata column."""

from __future__ import annotations

import json
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation


PROJECT_DIR = Path(__file__).resolve().parents[1]
BASE_TEMPLATE = PROJECT_DIR / "releases" / "plantilla_excel_v1_2" / "plantilla_plan_trabajo_jira_v1_2.xlsx"
HEADERS = [
    "Actividad *", "Área_Responsable", "Responsable *", "Tipo_de_seguimiento *",
    "Fecha_Objetivo", "Proyecto_o_Frente", "Contexto_o_Resultado", "Comentarios",
]


def build(catalog_path: Path, output_path: Path) -> None:
    document = json.loads(catalog_path.read_text(encoding="utf-8"))
    catalog = document.get("catalog", document)
    options = catalog.get("options") or []
    workbook = load_workbook(BASE_TEMPLATE)
    plan = workbook["PLAN_DE_TRABAJO"]
    catalogs = workbook["CATALOGOS"]

    for merged in list(plan.merged_cells.ranges):
        plan.unmerge_cells(str(merged))
    plan.insert_cols(2)
    plan.merge_cells("A1:H1")
    plan.merge_cells("A3:H3")
    plan["A1"] = "PLAN DE TRABAJO · Plantilla Jira v1.3"
    plan["A3"] = "Plantilla Jira v1.3 · Área responsable es un metadato local opcional"
    for column, value in enumerate(HEADERS, 1):
        plan.cell(5, column).value = value

    # Give the new area column the same input style as the neighboring text field.
    for row in range(5, 56):
        source = plan.cell(row, 3)
        target = plan.cell(row, 2)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.protection:
            target.protection = copy(source.protection)
    for column, width in {"A": 34, "B": 24, "C": 22, "D": 22, "E": 18, "F": 22, "G": 38, "H": 30}.items():
        plan.column_dimensions[column].width = width

    for validation in list(plan.data_validations.dataValidation):
        plan.data_validations.dataValidation.remove(validation)
    if options:
        validation = DataValidation(type="list", formula1=f"='CATALOGOS'!$A$2:$A${len(options) + 1}", allow_blank=True)
        plan.add_data_validation(validation)
        validation.add("D6:D55")

    for row in catalogs.iter_rows(min_row=1, max_row=max(catalogs.max_row, len(options) + 8), min_col=1, max_col=5):
        for cell in row:
            cell.value = None
    for row_index, values in enumerate([["display_value", "jira_option_id"]] + [
        [option.get("value"), option.get("jira_id")] for option in options
    ], 1):
        for column_index, value in enumerate(values, 1):
            catalogs.cell(row_index, column_index).value = value
    metadata = [
        ("metadata", "value"), ("template_version", "1.3"), ("space", catalog.get("space")),
        ("work_type", catalog.get("work_type")), ("field_id", catalog.get("field_id")),
        ("catalog_fetched_at", catalog.get("fetched_at")),
        ("catalog_status", catalog.get("status") or "unavailable"),
        ("area_responsable", "local metadata; not sent to Jira"),
    ]
    for row_index, values in enumerate(metadata, 1):
        catalogs.cell(row_index, 4).value, catalogs.cell(row_index, 5).value = values

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    manifest = {"version": "1.3", "sheet": "PLAN_DE_TRABAJO", "headers": HEADERS, "catalog": catalog,
                "area_responsable": {"source": "LOCAL", "required": False, "jira_field_id": None}}
    output_path.with_name("manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    output_path.with_name("README.md").write_text(
        "# Plantilla Excel Jira v1.3\n\n"
        "Agrega `Área_Responsable` como metadato local opcional. No se envía a Jira. "
        "Conserva los campos de Tipo de seguimiento de v1.2 y sigue aceptando archivos v1.2.\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    if len(sys.argv) not in (2, 3):
        raise SystemExit("Uso: build_template_v13.py CATALOG.json [OUTPUT.xlsx]")
    catalog_file = Path(sys.argv[1]).resolve()
    output_file = Path(sys.argv[2]).resolve() if len(sys.argv) == 3 else PROJECT_DIR / "releases" / "plantilla_excel_v1_3" / "plantilla_plan_trabajo_jira_v1_3.xlsx"
    build(catalog_file, output_file)
